"""
buat_laporan.py
Membuat laporan SPK (format: Cover, Kata Pengantar, Daftar Isi, BAB I-V) dalam .docx
Times New Roman 12, hitam. Semua angka dihitung langsung dari spk_core.py.

Pemakaian:
  python buat_laporan.py                      -> memakai DATA ILUSTRASI (uji alur)
  python buat_laporan.py --data folder_csv    -> memakai 12 file CSV Kaggle (hasil asli)
  python buat_laporan.py --data folder_csv --n 6   -> pakai 6 kolam pertama saja
"""
import argparse
import os
import re
import subprocess
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import FancyBboxPatch

import spk_core as core

FONT = "Times New Roman"
J, C, L, R_ = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT


# ------------------------------------------------------------------ util angka
def f(x, d=4):
    s = f"{x:.{d}f}".replace(".", ",")
    return s.replace("-", "−")


def fx(x):
    if float(x).is_integer():
        return str(int(x))
    return f"{x:g}".replace(".", ",")


# ------------------------------------------------------------------ util docx
def fix_style(st, bold=None):
    st.font.name = FONT
    st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.italic = False
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if bold is not None:
        st.font.bold = bold


def setup(doc):
    for st in doc.styles:
        try:
            if st.type in (1, 2, 3):  # paragraph, character, table
                if st.type == 1 or st.name in ("Default Paragraph Font",):
                    fix_style(st)
        except Exception:
            pass
    fix_style(doc.styles["Normal"], False)
    for i, name in enumerate(["Heading 1", "Heading 2", "Heading 3"], 1):
        st = doc.styles[name]
        fix_style(st, True)
        pf = st.paragraph_format
        pf.space_before = Pt(12 if i == 1 else 10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.line_spacing = 1.5
        pf.alignment = C if i == 1 else L
        if i == 3:
            pf.left_indent = Cm(0)
    # margin 4-3-3-3 cm (standar skripsi/proposal), A4
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(4), Cm(3), Cm(3), Cm(3)


def add_runs(p, text, bold=False, italic=False, hl=False):
    """Mendukung **tebal** inline."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


class Rpt:
    def __init__(self):
        self.doc = Document()
        setup(self.doc)
        self.headings = []  # (level, text)
        self.fig_no, self.tab_no = {}, {}

    # ---- teks
    def P(self, text, bold=False, align=J, first=True, after=6, ls=1.5, italic=False, hl=False,
          keep=False, left=0.0):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        pf.line_spacing = ls
        pf.space_after = Pt(after)
        pf.space_before = Pt(0)
        pf.keep_with_next = keep
        if left:
            pf.left_indent = Cm(left)
        if first and align == J:
            pf.first_line_indent = Cm(1.25)
        add_runs(p, text, bold=bold, italic=italic, hl=hl)
        return p

    def H(self, text, level):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        self.headings.append((level, text))
        return p

    def item(self, label, text, left=1.9, hang=0.63, after=3):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(left)
        pf.first_line_indent = Cm(-hang)
        pf.line_spacing = 1.5
        pf.space_after = Pt(after)
        pf.alignment = J
        pf.tab_stops.add_tab_stop(Cm(left))
        add_runs(p, f"{label}\t{text}")
        return p

    def bullets(self, lst, **kw):
        for i, t in enumerate(lst, 1):
            self.item(f"{i}.", t, **kw)

    def pagebreak(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- tabel
    def caption_tab(self, bab, text):
        self.tab_no[bab] = self.tab_no.get(bab, 0) + 1
        label = f"Tabel {bab}.{self.tab_no[bab]}"
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = C
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        add_runs(p, f"**{label}** {text}")
        return label

    def caption_fig(self, bab, text):
        self.fig_no[bab] = self.fig_no.get(bab, 0) + 1
        label = f"Gambar {bab}.{self.fig_no[bab]}"
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = C
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        add_runs(p, f"**{label}** {text}")
        return label

    def table(self, headers, rows, widths, center_cols=None, bold_first_col=False, font_bold_rows=()):
        center_cols = set(center_cols if center_cols is not None else range(len(headers)))
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, h in enumerate(headers):
            c = t.rows[0].cells[j]
            c.width = Cm(widths[j])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "D9D9D9")
            tcPr = c._tc.get_or_add_tcPr()
            va = tcPr.find(qn("w:vAlign"))
            if va is not None:
                va.addprevious(shd)
            else:
                tcPr.append(shd)
            p = c.paragraphs[0]
            p.paragraph_format.alignment = C
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, h, bold=True)
        # header berulang
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].width = Cm(widths[j])
                cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cells[j].paragraphs[0]
                p.paragraph_format.alignment = C if j in center_cols else L
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, str(v), bold=(bold_first_col and j == 0) or (i in font_bold_rows))
        # jangan pecah baris antar halaman
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trPr.append(cs)
        # lebar kolom eksplisit (grid + tblW) agar dibaca Word & LibreOffice
        for j, w in enumerate(widths):
            t.columns[j].width = Cm(w)
        tblPr = t._tbl.tblPr
        tw = tblPr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW")
            tblPr.append(tw)
        tw.set(qn("w:type"), "dxa")
        tw.set(qn("w:w"), str(int(sum(widths) * 567)))
        # tabel kecil: jaga agar tidak terpisah antar halaman
        if len(rows) <= 14:
            for row in t.rows[:-1]:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.paragraph_format.keep_with_next = True
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        return t

    def placeholder(self, label_text, height=5.0):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Cm(14)
        tr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(height * 567)))
        h.set(qn("w:hRule"), "atLeast")
        tr.append(h)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.alignment = C
        add_runs(p, label_text, bold=True, hl=True)

    def figure(self, path, width=13.5):
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = C
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width))

    def code(self, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Cm(14)
        first = True
        for line in text.strip("\n").split("\n"):
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.paragraph_format.alignment = L
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, line if line else " ")
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    def persamaan(self, text, no):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = L
        pf.line_spacing = 1.5
        pf.space_after = Pt(4)
        pf.tab_stops.add_tab_stop(Cm(7), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(14), WD_TAB_ALIGNMENT.RIGHT)
        add_runs(p, f"\t{text}\t({no})")


# ------------------------------------------------------------------ gambar otomatis
def gambar_alur(path):
    fig, ax = plt.subplots(figsize=(5.6, 8.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 20)
    ax.axis("off")
    plt.rcParams["font.family"] = "serif"
    plt.rcParams["font.serif"] = ["Liberation Serif", "Times New Roman", "DejaVu Serif"]

    def box(x, y, w, h, text, rounded=True, fc="#FFFFFF"):
        style = "round,pad=0.02,rounding_size=0.35" if rounded else "square,pad=0.02"
        ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h, boxstyle=style, ec="black", fc=fc, lw=1.1))
        ax.text(x, y, text, ha="center", va="center", fontsize=8.6, family="serif")

    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="-|>", color="black", lw=1.0))

    ys = [19, 16.9, 14.8, 12.7, 10.6]
    box(5, ys[0], 3.2, 1.0, "Mulai", True, "#E8E8E8")
    box(5, ys[1], 8.2, 1.4, "Pengumpulan dataset publik Kaggle\n(12 kolam, sensor IoT)", False)
    box(5, ys[2], 8.2, 1.4, "Pra-pemrosesan data\n(pembersihan, rata-rata per kolam)", False)
    box(5, ys[3], 8.2, 1.4, "Penentuan kriteria, tipe (Benefit/Cost),\nkonversi skala 1–5 dan bobot (W)", False)
    ax.add_patch(plt.Polygon([[5, 11.5], [8.6, 10.0], [5, 8.5], [1.4, 10.0]], fc="white", ec="black", lw=1.1))
    ax.text(5, 10.0, "Matriks keputusan\nX valid?", ha="center", va="center", fontsize=8.6)
    box(2.7, 6.6, 4.3, 1.5, "Perhitungan WP\nS = Π x^w\nV = S / ΣS", False)
    box(7.3, 6.6, 4.3, 1.5, "Perhitungan SAW\nR = normalisasi\nV = Σ w·R", False)
    box(5, 4.2, 8.2, 1.2, "Perangkingan alternatif (V terbesar)", False)
    box(5, 2.5, 8.2, 1.2, "Pengujian: Spearman, Kendall,\nanalisis sensitivitas", False)
    box(5, 0.9, 8.2, 1.1, "Keputusan: prioritas kolam", False)
    for a, b in zip(ys[:3], ys[1:4]):
        arrow(5, a - 0.55 if a == ys[0] else a - 0.72, 5, b + 0.72)
    arrow(5, ys[3] - 0.72, 5, 11.5)
    arrow(3.3, 9.3, 2.7, 7.4)
    arrow(6.7, 9.3, 7.3, 7.4)
    ax.text(1.9, 8.6, "Ya", fontsize=8.2)
    arrow(2.7, 5.85, 4.2, 4.8)
    arrow(7.3, 5.85, 5.8, 4.8)
    arrow(5, 3.6, 5, 3.1)
    arrow(5, 1.9, 5, 1.45)
    # jalur "Tidak"
    ax.annotate("", xy=(9.6, 14.8), xytext=(8.6, 10.0),
                arrowprops=dict(arrowstyle="-|>", color="black", lw=1.0,
                                connectionstyle="angle,angleA=0,angleB=90,rad=0"))
    ax.text(9.0, 10.3, "Tidak", fontsize=8.2)
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def gambar_hasil(h, path):
    alts = list(h["V_wp"].index)
    x = np.arange(len(alts))
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    plt.rcParams["font.family"] = "serif"
    ax.bar(x - 0.2, h["V_wp"].values, 0.4, label="WP", color="#4C6A92", edgecolor="black")
    ax.bar(x + 0.2, h["V_saw"].values / h["V_saw"].sum(), 0.4, label="SAW (dinormalkan ke jumlah 1)",
           color="#BFC7D1", edgecolor="black")
    ax.set_xticks(x)
    ax.set_xticklabels(alts, rotation=30, ha="right", fontsize=8)
    ax.set_ylabel("Nilai preferensi (V)", fontsize=9)
    ax.legend(fontsize=8)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=220)
    plt.close(fig)


# ------------------------------------------------------------------ isi laporan
def bangun(rp, h, demo, tmp, toc_pages=None):
    d = rp.doc
    K = core.KRITERIA
    raw, X, W, S, Vw, R, Vs = h["raw"], h["X"], h["W"], h["S"], h["V_wp"], h["R"], h["V_saw"]
    alts = list(raw.index)
    n = len(alts)
    rw, rs = h["rank_wp"], h["rank_saw"]
    kode = [k["kode"] for k in K]
    nama = {k["kode"]: k["nama"] for k in K}
    best_wp, best_saw = rw.idxmin(), rs.idxmin()
    worst_wp = rw.idxmax()

    def note_demo(txt):
        if demo:
            rp.P("CATATAN (hapus setelah memakai data Kaggle): " + txt, hl=True, first=False)

    # ================= COVER =================
    for t_, b in [("LAPORAN TUGAS", True), ("SISTEM PENDUKUNG KEPUTUSAN (SPK)", True)]:
        rp.P(t_, bold=b, align=C, first=False, after=6)
    rp.P("", first=False)
    rp.P("PERANGKINGAN KUALITAS AIR KOLAM BUDIDAYA PERIKANAN MENGGUNAKAN METODE "
         "WEIGHTED PRODUCT (WP) DAN SIMPLE ADDITIVE WEIGHTING (SAW)", bold=True, align=C, first=False)
    rp.P("(Studi Kasus: Pengembangan Sistem Monitoring IoT “Tambak Presisi”)", bold=True, align=C,
         first=False, after=18)
    rp.placeholder("[TEMPATKAN LOGO KAMPUS DI SINI]", height=4.0)
    rp.P("", first=False)
    rp.P("Disusun oleh:", bold=True, align=C, first=False, after=2)
    for line in ["[Nama Lengkap]", "[NIM]", "[Kelas]"]:
        rp.P(line, align=C, first=False, after=2, hl=True)
    rp.P("", first=False)
    rp.P("Dosen Pengampu:", bold=True, align=C, first=False, after=2)
    rp.P("[Nama Dosen Pengampu]", align=C, first=False, hl=True)
    rp.P("", first=False)
    for line in ["PROGRAM STUDI [NAMA PROGRAM STUDI]", "FAKULTAS [NAMA FAKULTAS]", "[NAMA UNIVERSITAS]", "2026"]:
        rp.P(line, bold=True, align=C, first=False, after=2)

    # ================= FRONT MATTER (section 2) =================
    sec2 = d.add_section(WD_SECTION.NEW_PAGE)
    # ---- Kata pengantar
    p = rp.P("KATA PENGANTAR", bold=True, align=C, first=False, after=10)
    rp.P("Puji syukur penulis panjatkan ke hadirat Tuhan Yang Maha Esa atas rahmat dan karunia-Nya sehingga "
         "laporan tugas Sistem Pendukung Keputusan (SPK) berjudul “Perangkingan Kualitas Air Kolam Budidaya "
         "Perikanan Menggunakan Metode Weighted Product (WP) dan Simple Additive Weighting (SAW)” ini dapat "
         "diselesaikan. Laporan ini disusun sebagai latihan penulisan karya ilmiah, khususnya struktur dan "
         "sistematika seminar proposal maupun tesis pada semester mendatang.")
    rp.P("Laporan ini membahas penerapan metode perangkingan multi-atribut (Multiple Attribute Decision Making/MADM) "
         "pada data kualitas air hasil sensor Internet of Things (IoT). Pembahasan mencakup profil dataset publik "
         "yang bersumber dari Kaggle, alur metodologi SPK, algoritma dan perhitungan manual, implementasi program, "
         "serta analisa hasil perhitungan menggunakan aplikasi yang dikembangkan.")
    rp.P("Penulis mengucapkan terima kasih kepada dosen pengampu mata kuliah yang telah memberikan arahan dan "
         "materi metode perangkingan, serta kepada semua pihak yang telah membantu penyusunan laporan ini. Penulis "
         "menyadari bahwa laporan ini masih memiliki kekurangan, sehingga kritik dan saran yang membangun sangat "
         "diharapkan demi perbaikan di masa mendatang.")
    rp.P("[Kota], [tanggal] 2026", align=R_, first=False, after=2, hl=True)
    rp.P("Penulis", align=R_, first=False)
    rp.pagebreak()
    # ---- Daftar isi (diisi nanti)
    rp.P("DAFTAR ISI", bold=True, align=C, first=False, after=10)
    toc_anchor = d.add_paragraph()  # akan diganti
    toc_anchor.paragraph_format.space_after = Pt(0)
    rp.toc_anchor = toc_anchor

    # ================= BODY (section 3) =================
    d.add_section(WD_SECTION.NEW_PAGE)

    # ---------------- BAB I
    rp.H("BAB I  PROFIL DATASET", 1)
    rp.H("1.1 Latar Belakang", 2)
    rp.P("Kualitas air merupakan faktor penentu keberhasilan budidaya perikanan. Suhu, kekeruhan, pH, dan kandungan "
         "oksigen terlarut yang berada di luar rentang optimum dapat menurunkan nafsu makan, memperlambat "
         "pertumbuhan, bahkan menyebabkan kematian massal. Pada proyek “Tambak Presisi”, parameter tersebut telah "
         "dipantau secara real-time menggunakan mikrokontroler ESP32 yang terhubung dengan sensor suhu DS18B20, "
         "sensor kekeruhan, sensor TDS, dan sensor pH. Data dikirim melalui protokol MQTT ke broker HiveMQ Cloud dan "
         "ditampilkan pada aplikasi mobile berbasis Flutter.")
    rp.P("Ketika pengelola memiliki lebih dari satu kolam, muncul persoalan baru: kolam mana yang kondisinya paling "
         "baik dan kolam mana yang paling mendesak untuk ditangani? Persoalan ini bersifat multikriteria karena "
         "setiap kolam dinilai dari beberapa parameter sekaligus dengan tingkat kepentingan yang berbeda. "
         "Metode Weighted Product (WP) dan Simple Additive Weighting (SAW) sebagai bagian dari MADM dipilih untuk "
         "menghasilkan peringkat kolam secara objektif dan dapat ditelusuri langkah perhitungannya.")
    rp.P("Karena jumlah kolam pada proyek “Tambak Presisi” masih terbatas, laporan ini memakai dataset publik dari "
         "Kaggle yang memuat data sensor IoT dari beberapa kolam budidaya sebagai bahan latihan. Pendekatan yang "
         "sama dapat langsung diterapkan pada data sensor proyek sendiri di tahap penelitian berikutnya.")
    rp.H("1.2 Tujuan dan Batasan", 2)
    rp.P("**Tujuan** laporan ini adalah sebagai berikut.", first=False, after=2)
    rp.bullets([
        "Menyusun profil dataset publik kualitas air kolam yang digunakan sebagai alternatif dan kriteria.",
        "Merancang alur metodologi SPK perangkingan menggunakan metode WP dan SAW.",
        "Menghitung secara manual dan terprogram nilai preferensi serta peringkat setiap kolam.",
        "Mengimplementasikan perhitungan ke dalam aplikasi berbasis Python dan menganalisis hasilnya.",
    ])
    rp.P("**Batasan** yang digunakan adalah sebagai berikut.", first=False, after=2)
    rp.bullets([
        "Metode yang digunakan hanya WP dan SAW (metode perangkingan); tidak digunakan metode peramalan "
        "maupun pembelajaran mesin sehingga tidak ada proses pelatihan (training) model.",
        "Alternatif adalah kolam, dan nilai kriteria setiap kolam diambil dari rata-rata pembacaan sensor.",
        "Bobot kriteria ditetapkan secara subjektif berdasarkan tingkat kepentingan agronomis (skala 1–5).",
    ])
    rp.H("1.3 Sumber dan Karakteristik Dataset", 2)
    rp.P("Dataset yang digunakan adalah **Sensor Based Aquaponics Fish Pond Datasets** yang dipublikasikan di Kaggle "
         "(https://www.kaggle.com/datasets/ogbuokiriblessing/sensor-based-aquaponics-fish-pond-datasets) dan "
         "dideskripsikan oleh Udanor dkk. (2022) pada jurnal Data in Brief. Dataset dikontribusikan ke komunitas "
         "riset di bawah lisensi Creative Commons Attribution 4.0 (CC BY 4.0). Data dikumpulkan menggunakan "
         "sistem IoT berbasis ESP-32 yang membaca enam parameter kualitas air pada 12 kolam ikan akuaponik dan "
         "mengirimkannya ke penyimpanan awan ThingSpeak setiap 5 detik. Dataset terdiri atas 12 berkas CSV, "
         "masing-masing mewakili satu kolam.")
    rp.P("Pemilihan dataset ini didasarkan pada tiga alasan. Pertama, setiap berkas mewakili satu kolam sehingga "
         "secara alami dapat dijadikan alternatif dalam perangkingan. Kedua, parameter suhu, kekeruhan, dan pH sama "
         "dengan sensor yang dipakai pada proyek “Tambak Presisi”. Ketiga, dataset bersifat terbuka dan dapat "
         "diunduh siapa pun sehingga hasil penelitian dapat direproduksi.")
    cap = rp.caption_tab(1, "Atribut dataset publik")
    rp.table(["No", "Atribut", "Keterangan", "Dipakai"], [
        [1, "created_at / entry_id", "Waktu dan nomor urut pembacaan", "Tidak"],
        [2, "Temperature", "Suhu air (°C)", "Ya (C1)"],
        [3, "Turbidity", "Kekeruhan air (NTU)", "Ya (C2)"],
        [4, "pH", "Derajat keasaman air", "Ya (C3)"],
        [5, "Dissolved Oxygen", "Oksigen terlarut dalam air", "Ya (C4)"],
        [6, "Ammonia, Nitrate", "Kadar amonia dan nitrat", "Tidak"],
        [7, "Fish length / weight", "Panjang dan berat ikan (pengukuran manual)", "Pembanding validasi"],
    ], [1.0, 4.0, 6.0, 3.0], center_cols=[0, 3])
    rp.P("Nama kolom pada berkas asli dapat sedikit berbeda dari tabel di atas; program yang dikembangkan mendeteksi "
         "kolom secara otomatis berdasarkan kata kunci (temp, turb, ph, oxygen).", first=False, italic=True)
    rp.H("1.4 Alternatif dan Kriteria", 2)
    rp.P(f"Alternatif yang dievaluasi berjumlah **{n} kolam** ({alts[0]} sampai {alts[-1]}). Kriteria yang digunakan "
         "ditunjukkan pada Tabel 1.2. Kriteria suhu dan pH memiliki rentang optimum (bukan semakin besar semakin baik), "
         "sehingga nilainya dikonversi lebih dahulu menjadi skor 1–5 pada Bab III. Kekeruhan bertipe Cost "
         "(semakin kecil semakin baik) dan oksigen terlarut bertipe Benefit (semakin besar semakin baik) dipakai "
         "dalam nilai aslinya, sehingga kedua tipe atribut yang diajarkan pada metode WP dan SAW sama-sama "
         "terwakili.")
    rp.caption_tab(1, "Kriteria, tipe, dan tingkat kepentingan")
    rp.table(["Kode", "Kriteria", "Satuan", "Tipe", "Kepentingan (1–5)", "Alasan"], [
        ["C1", "Suhu Air", "°C", "Benefit*", 4, "Memengaruhi metabolisme ikan"],
        ["C2", "Kekeruhan", "NTU", "Cost", 3, "Mengganggu penetrasi cahaya dan insang"],
        ["C3", "pH Air", "–", "Benefit*", 5, "Paling menentukan toksisitas dan kesehatan"],
        ["C4", "Oksigen Terlarut", "mg/L", "Benefit", 4, "Kebutuhan respirasi ikan"],
    ], [1.1, 2.9, 1.4, 1.7, 2.4, 4.5], center_cols=[0, 2, 3, 4])
    rp.P("*Setelah dikonversi menjadi skor 1–5, nilai skor bertipe Benefit.", first=False, italic=True)
    rp.P("Parameter TDS pada proyek “Tambak Presisi” tidak tersedia pada dataset ini sehingga digantikan oleh oksigen "
         "terlarut. Hal ini dicatat sebagai keterbatasan pada Bab V.")
    rp.H("1.5 Pra-pemrosesan Data", 2)
    rp.bullets([
        "Pembacaan seluruh berkas CSV dan penyeragaman nama kolom.",
        "Penghapusan baris yang memuat nilai kosong (missing value) pada kolom yang dipakai.",
        "Penghapusan nilai tidak fisik, misalnya pH di luar 0–14, suhu ≤ 0, dan kekeruhan atau oksigen terlarut negatif.",
        "Agregasi: rata-rata tiap parameter pada setiap kolam sehingga diperoleh satu baris per alternatif.",
    ])
    rp.H("1.6 Profil Statistik Data Hasil Agregasi", 2)
    rp.caption_tab(1, "Ringkasan statistik nilai kriteria antar kolam")
    rows = []
    for k in K:
        col = raw[k["nama"]]
        rows.append([k["kode"], k["nama"], f(col.min(), 2), f(col.max(), 2), f(col.mean(), 2), f(col.std(), 2)])
    rp.table(["Kode", "Kriteria", "Minimum", "Maksimum", "Rerata", "Simpangan Baku"], rows,
             [1.4, 4.2, 2.0, 2.2, 2.0, 2.2], center_cols=[0, 2, 3, 4, 5])
    note_demo("Seluruh angka pada laporan ini masih berasal dari DATA ILUSTRASI (bukan data Kaggle) agar alur "
              "perhitungan dapat diuji. Jalankan: python buat_laporan.py --data folder_csv_kaggle  untuk membuat "
              "ulang laporan dengan data asli.")
    rp.placeholder("[TEMPATKAN GAMBAR 1.1: Tangkapan layar halaman dataset di Kaggle\n"
                   "(judul, lisensi, daftar 12 berkas CSV)]", 5.5)
    rp.caption_fig(1, "Halaman dataset Sensor Based Aquaponics Fish Pond Datasets pada Kaggle")
    rp.placeholder("[TEMPATKAN GAMBAR 1.2: Hasil df.head() dan df.info() pada Python/Jupyter\n"
                   "untuk salah satu berkas CSV kolam]", 5.5)
    rp.caption_fig(1, "Cuplikan isi dan struktur data salah satu berkas CSV kolam")

    # ---------------- BAB II
    rp.pagebreak()
    rp.H("BAB II  ALUR PROSES BISNIS", 1)
    rp.H("2.1 Gambaran Proses Bisnis", 2)
    rp.P("Proses bisnis yang didukung adalah penentuan prioritas penanganan kolam oleh pengelola tambak. "
         "Masukan sistem berupa data kualitas air tiap kolam, proses berupa perhitungan nilai preferensi "
         "menggunakan metode WP dan SAW, dan keluaran berupa peringkat kolam beserta rekomendasi prioritas. "
         "Pengelola berperan sebagai pengambil keputusan, sedangkan sistem berperan menyajikan hasil "
         "perhitungan yang objektif dan dapat ditelusuri.")
    rp.H("2.2 Landasan Metode Perangkingan", 2)
    rp.H("2.2.1 Multiple Attribute Decision Making (MADM)", 3)
    rp.P("MADM adalah kelompok metode Sistem Pendukung Keputusan yang memilih atau mengurutkan alternatif "
         "terbaik dari sejumlah alternatif berdasarkan banyak atribut (kriteria). Setiap kriteria memiliki bobot "
         "dan sifat tertentu, yaitu Benefit (semakin besar semakin baik) atau Cost (semakin kecil semakin baik).")
    rp.H("2.2.2 Weighted Product (WP)", 3)
    rp.P("Metode WP menggunakan perkalian untuk menghubungkan rating atribut, dengan rating setiap atribut "
         "dipangkatkan dengan bobot atribut yang bersangkutan. Tahapan perhitungannya adalah sebagai berikut.",
         first=False, after=2)
    rp.item("1.", "**Menentukan kriteria (C)** beserta sifatnya (Benefit atau Cost).")
    rp.item("2.", "**Menentukan bobot (W)** setiap kriteria, lalu memperbaikinya sehingga ΣW = 1:")
    rp.persamaan("w_j = W_j / Σ W_j", "2.1")
    rp.item("3.", "**Menentukan vektor S** (hasil normalisasi) dengan pangkat positif untuk Benefit dan pangkat "
                  "negatif untuk Cost:")
    rp.persamaan("S_i = Π (x_ij) ^ (± w_j),  j = 1 … n", "2.2")
    rp.item("4.", "**Menentukan vektor V** (nilai preferensi) sebagai proporsi S setiap alternatif:")
    rp.persamaan("V_i = S_i / Σ S_i", "2.3")
    rp.H("2.2.3 Simple Additive Weighting (SAW)", 3)
    rp.P("Metode SAW dikenal sebagai metode penjumlahan terbobot. Konsep dasarnya adalah mencari penjumlahan "
         "terbobot dari rating kinerja pada setiap alternatif untuk semua atribut. Matriks keputusan dinormalisasi "
         "ke skala yang dapat diperbandingkan:", first=False, after=2)
    rp.persamaan("r_ij = x_ij / max_i(x_ij)   (Benefit)", "2.4")
    rp.persamaan("r_ij = min_i(x_ij) / x_ij   (Cost)", "2.5")
    rp.P("Nilai preferensi setiap alternatif dihitung dengan:", first=False, after=2)
    rp.persamaan("V_i = Σ ( w_j · r_ij )", "2.6")
    rp.P("Alternatif dengan V terbesar merupakan alternatif terbaik. Perbedaan utama kedua metode: WP memakai "
         "perkalian sehingga alternatif yang sangat buruk pada satu kriteria terpenalti lebih kuat, sedangkan SAW "
         "memakai penjumlahan sehingga kekurangan pada satu kriteria dapat ditutup oleh kelebihan pada kriteria lain.")
    rp.H("2.3 Alur Metodologi SPK yang Dibuat", 2)
    p1 = os.path.join(tmp, "alur.png")
    gambar_alur(p1)
    rp.figure(p1, 11.0)
    rp.caption_fig(2, "Alur metodologi SPK perangkingan kolam dengan metode WP dan SAW")
    rp.P("Penjelasan setiap tahap alur pada Gambar 2.1 adalah sebagai berikut.", first=False, after=2)
    rp.item("1.", "**Pengumpulan dataset.** Mengunduh 12 berkas CSV dari Kaggle; setiap berkas menjadi satu alternatif.")
    rp.item("2.", "**Pra-pemrosesan.** Membersihkan nilai kosong dan tidak fisik, lalu menghitung rata-rata tiap "
                  "parameter per kolam.")
    rp.item("3.", "**Penentuan kriteria dan bobot.** Menetapkan C1–C4, tipe Benefit/Cost, mengonversi suhu dan pH ke "
                  "skor 1–5, dan menormalkan bobot hingga ΣW = 1.")
    rp.item("4.", "**Validasi matriks keputusan.** Memastikan tidak ada nilai kosong atau nol pada kriteria bertipe Cost "
                  "yang dapat menyebabkan pembagian dengan nol.")
    rp.item("5.", "**Perhitungan WP dan SAW secara paralel** menggunakan matriks dan bobot yang sama.")
    rp.item("6.", "**Perangkingan.** Mengurutkan kolam berdasarkan V dari terbesar ke terkecil.")
    rp.item("7.", "**Pengujian hasil.** Membandingkan peringkat WP dan SAW (Spearman, Kendall) serta menguji "
                  "kestabilan dengan analisis sensitivitas bobot.")
    rp.item("8.", "**Keputusan.** Menyajikan peringkat dan kolam yang perlu diprioritaskan penanganannya.")
    rp.H("2.4 Keterkaitan dengan Sistem IoT Tambak Presisi", 2)
    rp.P("SPK ini berfungsi sebagai lapisan keputusan di atas sistem monitoring yang sudah dibangun. Data sensor "
         "ESP32 yang dikirim melalui MQTT dapat dirata-ratakan per kolam (misalnya per hari) dan menjadi masukan "
         "matriks keputusan. Pada laporan ini, data publik dipakai sebagai pengganti sementara untuk memvalidasi alur "
         "metodologi.")
    rp.placeholder("[TEMPATKAN GAMBAR 2.2: Diagram arsitektur sistem Tambak Presisi\n"
                   "(Sensor → ESP32 → HiveMQ MQTT → Aplikasi Flutter → modul SPK)]", 6.0)
    rp.caption_fig(2, "Arsitektur sistem IoT Tambak Presisi dan posisi modul SPK")

    # ---------------- BAB III
    rp.pagebreak()
    rp.H("BAB III  ALGORITMA PERHITUNGAN", 1)
    rp.H("3.1 Alur Perhitungan", 2)
    rp.P("Algoritma perhitungan bersifat deterministik dan tidak memerlukan pelatihan model. Langkah-langkahnya "
         "adalah sebagai berikut.", first=False, after=2)
    rp.bullets([
        "Bentuk matriks keputusan X (alternatif × kriteria) dari rata-rata data kolam.",
        "Konversi nilai suhu dan pH menjadi skor 1–5 sesuai aturan pada Tabel 3.2.",
        "Hitung bobot ternormalisasi w_j = W_j / ΣW_j.",
        "WP: hitung S_i = Π x_ij^(±w_j), kemudian V_i = S_i / ΣS_i.",
        "SAW: normalisasi r_ij, kemudian V_i = Σ w_j · r_ij.",
        "Urutkan V menurun untuk memperoleh peringkat, lalu hitung korelasi peringkat dan uji sensitivitas.",
    ])
    rp.H("3.2 Data Masukan (Matriks Keputusan Awal)", 2)
    note_demo("Tabel 3.1 berisi data ilustrasi. Ganti dengan rata-rata data Kaggle dengan menjalankan program.")
    rp.caption_tab(3, "Rata-rata nilai kriteria tiap kolam (data awal)")
    rows = [[a] + [f(raw.loc[a, k["nama"]], 2) for k in K] for a in alts]
    rp.table(["Alternatif"] + [f"{k['kode']} ({k['satuan']})" if k["satuan"] else k["kode"] for k in K], rows,
             [3.0, 2.5, 2.5, 2.5, 2.5], center_cols=[1, 2, 3, 4])
    rp.H("3.3 Konversi Nilai ke Skala 1–5", 2)
    rp.P("Suhu dan pH memiliki rentang optimum sehingga nilai yang jauh di bawah maupun di atas rentang sama-sama "
         "buruk. Nilai dalam rentang optimum diberi skor 5; setiap pergeseran satu langkah dari rentang optimum "
         "menurunkan skor satu tingkat. Rentang optimum suhu 28–30 °C (langkah 2 °C) dan pH 7,0–8,0 (langkah 0,5) "
         "mengacu pada rentang aman yang dipakai pada firmware dan literatur budidaya ikan tropis (suhu 25–32 °C, "
         "pH 6,5–8,5).")
    rp.caption_tab(3, "Aturan konversi skor kriteria C1 dan C3")
    rows = []
    a1, a3 = core.ATURAN_SKOR["C1"], core.ATURAN_SKOR["C3"]

    def rentang(a, s):
        lo, hi = a["inti"]
        st = a["langkah"]
        if s == 5:
            return f"{fx(lo)} ≤ x ≤ {fx(hi)}"
        j = 5 - s
        if s == 1:
            return f"x < {fx(lo - 3 * st)} atau x > {fx(hi + 3 * st)}"
        return (f"{fx(lo - j * st)} ≤ x < {fx(lo - (j - 1) * st)} atau "
                f"{fx(hi + (j - 1) * st)} < x ≤ {fx(hi + j * st)}")
    for s in (5, 4, 3, 2, 1):
        rows.append([s, rentang(a1, s), rentang(a3, s)])
    rp.table(["Skor", "C1 Suhu (°C)", "C3 pH"], rows, [1.5, 6.5, 6.0], center_cols=[0])
    rp.caption_tab(3, "Matriks keputusan X setelah konversi")
    rows = [[a, f"{int(X.loc[a, 'C1'])}", f(X.loc[a, "C2"], 2), f"{int(X.loc[a, 'C3'])}", f(X.loc[a, "C4"], 2)]
            for a in alts]
    rp.table(["Alternatif", "C1 (skor)", "C2 (NTU)", "C3 (skor)", "C4 (mg/L)"], rows,
             [3.0, 2.5, 2.5, 2.5, 2.5], center_cols=[1, 2, 3, 4])
    rp.H("3.4 Penentuan Bobot", 2)
    kep = {k["kode"]: k["kepentingan"] for k in K}
    rp.P(f"Tingkat kepentingan kriteria dijumlahkan (ΣW = {sum(kep.values())}) dan dinormalkan "
         "dengan persamaan 2.1.")
    rp.caption_tab(3, "Bobot kriteria sebelum dan sesudah normalisasi")
    rows = [[k["kode"], k["nama"], k["tipe"], k["kepentingan"],
             f"{k['kepentingan']} / {sum(kep.values())}", f(W[k["kode"]], 4)] for k in K]
    rows.append(["", "Jumlah", "", sum(kep.values()), "", f(W.sum(), 4)])
    rp.table(["Kode", "Kriteria", "Tipe", "W", "Perhitungan", "w (normal)"], rows,
             [1.4, 4.2, 1.8, 1.2, 2.8, 2.6], center_cols=[0, 2, 3, 4, 5], font_bold_rows=(len(rows) - 1,))
    rp.H("3.5 Perhitungan Metode Weighted Product (WP)", 2)
    rp.H("3.5.1 Pangkat Bobot", 3)
    rp.P("Pangkat bernilai positif untuk kriteria Benefit dan negatif untuk Cost (persamaan 2.2). "
         "Pada kasus ini hanya C2 (Kekeruhan) yang berpangkat negatif.")
    rp.caption_tab(3, "Pangkat bobot pada metode WP")
    rows = [[k["kode"], k["tipe"], f(h["pangkat"][k["kode"]], 4)] for k in K]
    rp.table(["Kriteria", "Tipe", "Pangkat"], rows, [3.0, 3.0, 3.0], center_cols=[0, 1, 2])
    rp.H("3.5.2 Perhitungan Vektor S", 3)
    rp.P("Nilai S dihitung dengan mengalikan seluruh x_ij yang dipangkatkan sesuai Tabel 3.5.")
    rp.caption_tab(3, "Perhitungan vektor S")
    rows = []
    for i, a in enumerate(alts, 1):
        parts = " × ".join(f"({fx(X.loc[a, c])}^{f(h['pangkat'][c], 4)})" for c in kode)
        rows.append([f"S{i} ({a})", parts, f(S[a], 4)])
    rp.table(["Alternatif", "Perhitungan", "S"], rows, [2.8, 8.4, 2.4], center_cols=[0, 2])
    rp.H("3.5.3 Perhitungan Vektor V dan Peringkat", 3)
    rp.P(f"Jumlah seluruh S adalah ΣS = {f(S.sum(), 4)}. Nilai V diperoleh dari persamaan 2.3, misalnya "
         f"V({alts[0]}) = {f(S[alts[0]], 4)} / {f(S.sum(), 4)} = {f(Vw[alts[0]], 4)}.")
    rp.caption_tab(3, "Nilai preferensi dan peringkat metode WP")
    rows = [[a, f(S[a], 4), f(Vw[a], 4), int(rw[a])] for a in alts]
    rows.append(["Jumlah", f(S.sum(), 4), f(Vw.sum(), 4), ""])
    rp.table(["Alternatif", "S", "V", "Peringkat"], rows, [3.5, 3.0, 3.0, 2.5], center_cols=[1, 2, 3],
             font_bold_rows=(len(rows) - 1,))
    rp.H("3.6 Perhitungan Metode Simple Additive Weighting (SAW)", 2)
    rp.H("3.6.1 Normalisasi Matriks", 3)
    mx = {c: X[c].max() for c in kode}
    mn = {c: X[c].min() for c in kode}
    rp.P("Nilai maksimum untuk kriteria Benefit dan minimum untuk kriteria Cost: "
         + "; ".join(f"{c} = {fx(mx[c] if K[kode.index(c)]['tipe'] == 'Benefit' else mn[c])}" for c in kode)
         + f". Contoh {alts[0]}, C2 (Cost): r = {fx(mn['C2'])} / {fx(X.loc[alts[0], 'C2'])} = "
         f"{f(R.loc[alts[0], 'C2'], 4)}.")
    rp.caption_tab(3, "Matriks ternormalisasi R")
    rows = [[a] + [f(R.loc[a, c], 4) for c in kode] for a in alts]
    rp.table(["Alternatif"] + kode, rows, [3.0, 2.4, 2.4, 2.4, 2.4], center_cols=[1, 2, 3, 4])
    rp.H("3.6.2 Perhitungan Nilai Preferensi", 3)
    a0 = alts[0]
    rp.P(f"Contoh perhitungan {a0}: V = " + " + ".join(f"({f(W[c], 4)} × {f(R.loc[a0, c], 4)})" for c in kode)
         + f" = {f(Vs[a0], 4)}.")
    rp.caption_tab(3, "Nilai preferensi dan peringkat metode SAW")
    rows = [[a] + [f(W[c] * R.loc[a, c], 4) for c in kode] + [f(Vs[a], 4), int(rs[a])] for a in alts]
    rp.table(["Alternatif"] + [f"w×{c}" for c in kode] + ["V", "Rank"], rows,
             [2.6, 1.7, 1.7, 1.7, 1.7, 2.0, 1.5], center_cols=[1, 2, 3, 4, 5, 6])
    rp.H("3.7 Perbandingan Hasil WP dan SAW", 2)
    rp.caption_tab(3, "Perbandingan peringkat WP dan SAW")
    rows = [[a, f(Vw[a], 4), int(rw[a]), f(Vs[a], 4), int(rs[a]), int(rw[a] - rs[a]), int((rw[a] - rs[a]) ** 2)]
            for a in alts]
    rows.append(["Jumlah", "", "", "", "", "", int(((rw - rs) ** 2).sum())])
    rp.table(["Alternatif", "V WP", "Rank WP", "V SAW", "Rank SAW", "d", "d²"], rows,
             [2.7, 1.9, 1.8, 1.9, 1.9, 1.3, 1.3], center_cols=[1, 2, 3, 4, 5, 6], font_bold_rows=(len(rows) - 1,))
    rp.H("3.8 Rumus Perhitungan pada Python", 2)
    rp.P("Seluruh perhitungan manual di atas diimplementasikan pada Python menggunakan pandas, numpy, dan scipy "
         "(berkas spk_core.py). Inti perhitungannya adalah sebagai berikut.", first=False)
    rp.code('''import pandas as pd
from scipy.stats import spearmanr, kendalltau

# X: DataFrame matriks keputusan (baris = alternatif, kolom = C1..C4)
# kepentingan: dict {'C1': 4, ...}   tipe: dict {'C1': 'Benefit', ...}
W = pd.Series(kepentingan) / sum(kepentingan.values())      # bobot normal

# --- Weighted Product ---
pangkat = pd.Series({c: (W[c] if tipe[c] == 'Benefit' else -W[c])
                     for c in X.columns})
S = (X ** pangkat).prod(axis=1)                             # vektor S
V_wp = S / S.sum()                                          # vektor V

# --- Simple Additive Weighting ---
R = pd.DataFrame({c: (X[c] / X[c].max() if tipe[c] == 'Benefit'
                      else X[c].min() / X[c]) for c in X.columns})
V_saw = (R * W).sum(axis=1)

# --- Perangkingan ---
rank_wp  = V_wp.rank(ascending=False, method='min').astype(int)
rank_saw = V_saw.rank(ascending=False, method='min').astype(int)

# --- Pengujian hasil ---
rho, _ = spearmanr(V_wp, V_saw)
tau, _ = kendalltau(V_wp, V_saw)''')
    rp.H("3.9 Pengukuran Kesalahan dan Pengujian Hasil", 2)
    rp.P("Metode WP dan SAW adalah metode perangkingan deterministik. Tidak ada model yang dilatih dan tidak ada "
         "variabel target hasil prediksi, sehingga ukuran kesalahan peramalan seperti MAE, MSE, atau RMSE tidak "
         "relevan. Sebagai gantinya, kualitas hasil perangkingan diuji dengan tiga cara berikut.")
    rp.P("**(1) Korelasi peringkat antar metode.** Koefisien Spearman dihitung dengan:", first=False, after=2)
    rp.persamaan("ρ = 1 − [ 6 Σ d² ] / [ n (n² − 1) ]", "3.1")
    sd2 = int(((rw - rs) ** 2).sum())
    ties = rw.nunique() < n or rs.nunique() < n
    rho_man = 1 - 6 * sd2 / (n * (n * n - 1))
    if not ties:
        rp.P(f"Dengan Σd² = {sd2} dan n = {n} (Tabel 3.10), diperoleh ρ = 1 − (6 × {sd2}) / ({n} × ({n}² − 1)) = "
             f"{f(rho_man, 4)}. Koefisien Kendall Tau dari program: τ = {f(h['tau'], 4)}.")
    else:
        rp.P(f"Karena terdapat peringkat kembar, korelasi dihitung program: ρ = {f(h['rho'], 4)} dan "
             f"τ = {f(h['tau'], 4)}.")
    rp.P("Nilai mendekati 1 berarti kedua metode menghasilkan urutan yang konsisten.")
    rp.P("**(2) Analisis sensitivitas bobot.** Tingkat kepentingan setiap kriteria dinaikkan dan diturunkan satu "
         "tingkat (batas 1–5), kemudian peringkat dihitung ulang. Hasilnya disajikan pada Bab V.")
    rp.P("**(3) Validasi eksternal (opsional).** Karena dataset memuat berat ikan hasil pengukuran manual, rata-rata "
         "berat ikan per kolam dapat dipakai sebagai pembanding. Korelasi Spearman antara V dan rata-rata berat ikan "
         "yang tinggi menunjukkan bahwa kolam berperingkat baik juga mendukung pertumbuhan ikan yang baik. "
         "Validasi ini hanya dilakukan bila kolom berat ikan tersedia pada berkas CSV.")

    # ---------------- BAB IV
    rp.pagebreak()
    rp.H("BAB IV  IMPLEMENTASI PROGRAM", 1)
    rp.H("4.1 Lingkungan Pengembangan", 2)
    rp.P("Aplikasi dikembangkan menggunakan bahasa Python 3 dengan pustaka pandas dan numpy (pengolahan data), "
         "scipy (korelasi peringkat), matplotlib (grafik), dan Streamlit (antarmuka pengguna berbasis web). "
         "Aplikasi terdiri atas dua berkas utama, yaitu spk_core.py yang memuat seluruh rumus perhitungan dan "
         "app_spk.py yang memuat antarmuka. Aplikasi dijalankan dengan perintah: streamlit run app_spk.py.")
    rp.H("4.2 Rancangan Struktur Antarmuka", 2)
    rp.P("Antarmuka dirancang dengan satu menu navigasi di sisi kiri (sidebar) yang memuat enam form, diurutkan "
         "mengikuti alur metodologi pada Bab II. Pengguna bergerak dari atas ke bawah: memasukkan data, menetapkan "
         "bobot, melihat hasil WP dan SAW, lalu membandingkan hasilnya.")
    rp.caption_tab(4, "Daftar form aplikasi")
    rp.table(["No", "Form", "Fungsi"], [
        [1, "Beranda", "Ringkasan tujuan dan petunjuk penggunaan"],
        [2, "Input Data", "Memuat berkas CSV kolam atau data contoh"],
        [3, "Kriteria dan Bobot", "Mengatur tingkat kepentingan dan melihat bobot normal"],
        [4, "Hasil WP", "Menampilkan S, V, dan peringkat metode WP"],
        [5, "Hasil SAW", "Menampilkan matriks R, V, dan peringkat metode SAW"],
        [6, "Perbandingan dan Validasi", "Membandingkan WP–SAW, korelasi, dan sensitivitas"],
    ], [1.0, 4.5, 8.5], center_cols=[0])
    forms = [
        ("4.3 Form Beranda",
         "Form Beranda adalah halaman pertama yang tampil ketika aplikasi dibuka. Form ini memuat judul aplikasi, "
         "penjelasan singkat tujuan SPK, dan langkah penggunaan enam form. Tidak ada masukan data pada form ini; "
         "fungsinya sebagai orientasi pengguna.",
         "Tampilan form Beranda"),
        ("4.4 Form Input Data",
         "Form Input Data berfungsi memasukkan alternatif. Pengguna dapat memilih salah satu dari dua sumber: "
         "(a) mengunggah beberapa berkas CSV hasil unduhan Kaggle sekaligus, di mana setiap berkas dianggap satu "
         "kolam, atau (b) memakai data contoh. Setelah data dimuat, aplikasi mendeteksi kolom suhu, kekeruhan, pH, "
         "dan oksigen terlarut, membersihkan data, dan menampilkan tabel rata-rata per kolam sebagai matriks "
         "keputusan awal.",
         "Tampilan form Input Data beserta tabel rata-rata per kolam"),
        ("4.5 Form Kriteria dan Bobot",
         "Form ini menampilkan empat kriteria beserta tipe Benefit atau Cost. Pengguna mengatur tingkat "
         "kepentingan setiap kriteria dengan penggeser (slider) skala 1–5. Aplikasi menghitung bobot "
         "ternormalisasi secara langsung dan menampilkan jumlahnya (harus 1). Perubahan pada form ini otomatis "
         "memengaruhi hasil pada form berikutnya.",
         "Tampilan form Kriteria dan Bobot"),
        ("4.6 Form Hasil WP",
         "Form Hasil WP menampilkan secara berurutan matriks keputusan setelah konversi skor, pangkat bobot, "
         "vektor S, vektor V, dan peringkat. Hasil diurutkan dari peringkat terbaik dan dilengkapi diagram batang "
         "nilai V sehingga pengguna dapat menelusuri setiap langkah perhitungan manual pada Bab III.",
         "Tampilan form Hasil WP"),
        ("4.7 Form Hasil SAW",
         "Form Hasil SAW menampilkan matriks ternormalisasi R, kontribusi tiap kriteria (w × r), nilai V, dan "
         "peringkat metode SAW. Struktur tampilannya dibuat sejajar dengan form Hasil WP agar mudah dibandingkan.",
         "Tampilan form Hasil SAW"),
        ("4.8 Form Perbandingan dan Validasi",
         "Form terakhir memuat tabel peringkat WP dan SAW berdampingan, koefisien Spearman dan Kendall, "
         "tabel analisis sensitivitas bobot, serta (jika kolom berat ikan tersedia) korelasi hasil perangkingan "
         "dengan rata-rata berat ikan. Form ini menjadi dasar kesimpulan dan rekomendasi keputusan.",
         "Tampilan form Perbandingan dan Validasi"),
    ]
    for i, (jdl, narasi, cap_) in enumerate(forms, 1):
        rp.H(jdl, 2)
        rp.P(narasi)
        rp.placeholder(f"[TEMPATKAN GAMBAR 4.{i}: Screenshot {cap_.lower()}]", 6.0)
        rp.caption_fig(4, cap_)
    rp.H("4.9 Cara Menjalankan Program", 2)
    rp.code('''pip install pandas numpy scipy matplotlib streamlit python-docx
streamlit run app_spk.py                       # menjalankan aplikasi SPK
python buat_laporan.py --data folder_csv       # membuat ulang laporan dari data Kaggle''')

    # ---------------- BAB V
    rp.pagebreak()
    rp.H("BAB V  HASIL DAN PEMBAHASAN", 1)
    rp.H("5.1 Luaran Perhitungan Aplikasi", 2)
    note_demo("Seluruh nilai pada bab ini berasal dari data ilustrasi. Setelah laporan dibuat ulang dengan data "
              "Kaggle, tabel, grafik, dan analisis pada bab ini otomatis berubah mengikuti data asli.")
    rp.P(f"Aplikasi menghitung nilai preferensi dan peringkat {n} kolam. Hasil akhir kedua metode disajikan pada "
         "Tabel 5.1, sedangkan Gambar 5.1 menampilkan keluaran halaman hasil pada aplikasi.")
    rp.caption_tab(5, "Hasil akhir perangkingan kolam")
    order = Vw.sort_values(ascending=False).index
    rows = [[a, f(Vw[a], 4), int(rw[a]), f(Vs[a], 4), int(rs[a])] for a in order]
    rp.table(["Alternatif", "V WP", "Peringkat WP", "V SAW", "Peringkat SAW"], rows,
             [3.2, 2.5, 2.8, 2.5, 3.0], center_cols=[1, 2, 3, 4])
    rp.placeholder("[TEMPATKAN GAMBAR 5.1: Screenshot hasil pada aplikasi (form Hasil WP, Hasil SAW,\n"
                   "dan form Perbandingan)]", 6.0)
    rp.caption_fig(5, "Luaran hasil perhitungan pada aplikasi SPK")
    p2 = os.path.join(tmp, "hasil.png")
    gambar_hasil(h, p2)
    rp.figure(p2, 13.0)
    rp.caption_fig(5, "Perbandingan nilai preferensi WP dan SAW tiap kolam")
    rp.H("5.2 Analisa Hasil Metode WP", 2)
    sel = Vw.sort_values(ascending=False)
    rp.P(f"Pada metode WP, **{best_wp}** menempati peringkat 1 dengan V = {f(Vw[best_wp], 4)}, sedangkan "
         f"**{worst_wp}** berada di peringkat terakhir dengan V = {f(Vw[worst_wp], 4)}. Rasio nilai V tertinggi "
         f"terhadap terendah sebesar {f(sel.iloc[0] / sel.iloc[-1], 2)} kali, yang menunjukkan bahwa perbedaan kualitas "
         "antar kolam cukup jelas. Karena WP berbasis perkalian, kolam yang buruk pada satu kriteria berbobot "
         "besar (misalnya pH, bobot " + f(W['C3'], 4) + ") terpenalti kuat pada nilai S.")
    rp.H("5.3 Analisa Hasil Metode SAW", 2)
    lemah = {c: R.loc[worst_wp, c] for c in kode}
    c_lemah = min(lemah, key=lemah.get)
    rp.P(f"Pada metode SAW, **{best_saw}** menempati peringkat 1 dengan V = {f(Vs[best_saw], 4)}. "
         f"{worst_wp} memperoleh nilai ternormalisasi terendah pada kriteria {c_lemah} ({nama[c_lemah]}) "
         f"sebesar {f(lemah[c_lemah], 4)}, sehingga kriteria itulah yang paling perlu diperbaiki pada kolam tersebut.")
    rp.H("5.4 Perbandingan dan Pengujian Hasil", 2)
    top_txt = ("sama" if h["sama_top1"] else "berbeda")
    rp.P(f"Korelasi Spearman antara peringkat WP dan SAW adalah ρ = {f(h['rho'], 4)} dan Kendall τ = "
         f"{f(h['tau'], 4)}. Kolam peringkat 1 pada kedua metode {top_txt}, dan total selisih peringkat mutlak "
         f"sebesar {h['selisih_rank']}. "
         + ("Nilai korelasi yang tinggi menunjukkan bahwa kedua metode konsisten sehingga keputusan tidak "
            "bergantung pada pemilihan metode." if h["rho"] >= 0.8 else
            "Nilai korelasi yang tidak terlalu tinggi menunjukkan perbedaan karakter kedua metode yang perlu dikaji "
            "lebih lanjut."))
    rp.P("Hasil analisis sensitivitas bobot disajikan pada Tabel 5.2.")
    rp.caption_tab(5, "Analisis sensitivitas bobot")
    sens = h["sensitivitas"]
    rows = [[r["Skenario"], r["Peringkat 1 WP"], r["Peringkat 1 SAW"], f(r["Spearman WP vs dasar"], 4)]
            for _, r in sens.iterrows()]
    rp.table(["Skenario", "Peringkat 1 WP", "Peringkat 1 SAW", "Spearman vs dasar"], rows,
             [2.8, 3.6, 3.6, 3.6], center_cols=[0, 1, 2, 3])
    stabil = (sens["Peringkat 1 WP"] == best_wp).all() and (sens["Peringkat 1 SAW"] == best_saw).all()
    rp.P("Pada seluruh skenario perubahan bobot, kolam peringkat 1 " +
         ("tidak berubah, sehingga hasil perangkingan dinilai stabil terhadap perubahan bobot yang wajar."
          if stabil else
          "berubah pada sebagian skenario, sehingga hasil perangkingan sensitif terhadap penetapan bobot "
          "dan perlu pembobotan yang lebih objektif (misalnya AHP)."))
    if "valid_rho_wp" in h:
        rp.P(f"Pada validasi eksternal, korelasi Spearman antara V dengan rata-rata berat ikan sebesar "
             f"{f(h['valid_rho_wp'], 4)} (WP) dan {f(h['valid_rho_saw'], 4)} (SAW).")
    else:
        rp.P("Validasi eksternal dengan berat ikan belum dapat ditampilkan karena kolom berat ikan tidak tersedia "
             "pada data yang dimuat.", hl=demo)
    rp.H("5.5 Implikasi bagi Pengelola Tambak", 2)
    rp.P(f"Berdasarkan peringkat, pengelola dapat memprioritaskan penanganan pada {worst_wp} dan kolam berperingkat "
         f"rendah lainnya, sementara {best_wp} dapat dijadikan acuan kondisi ideal. Untuk kolam berperingkat rendah, "
         "kriteria dengan nilai ternormalisasi terkecil menjadi titik awal perbaikan, misalnya penggantian air "
         "untuk menurunkan kekeruhan, aerasi untuk menaikkan oksigen terlarut, dan penyesuaian pH.")
    rp.H("5.6 Keterbatasan dan Saran", 2)
    rp.bullets([
        "Dataset tidak memuat TDS sehingga tidak sepenuhnya sama dengan sensor proyek Tambak Presisi; "
        "pada penelitian lanjutan, data sensor sendiri (termasuk TDS) dapat dipakai.",
        "Bobot ditetapkan subjektif; penelitian lanjutan dapat memakai AHP atau entropi untuk bobot yang lebih objektif.",
        "Konversi skor memakai rentang optimum umum; rentang spesifik spesies (misalnya nila atau udang vaname) "
        "sebaiknya dikonsultasikan dengan pakar budidaya.",
        "Data diagregasi menjadi rata-rata sehingga dinamika waktu tidak tertangkap; peramalan deret waktu "
        "(misalnya dengan jaringan saraf tiruan) dapat menjadi topik lanjutan pada tesis atau proposal berikutnya.",
    ])
    rp.H("5.7 Kesimpulan Sementara", 2)
    rp.P(f"Metode WP dan SAW berhasil menghasilkan peringkat {n} kolam secara terukur. Aplikasi yang dibuat mampu "
         "menampilkan seluruh langkah perhitungan yang sama dengan perhitungan manual, dan pengujian korelasi serta "
         "sensitivitas menunjukkan tingkat konsistensi hasil sebagaimana dijelaskan pada subbab 5.4.")

    # ---------------- DAFTAR PUSTAKA
    rp.pagebreak()
    rp.P("DAFTAR PUSTAKA", bold=True, align=C, first=False, after=10)
    refs = [
        "Udanor, C. N., Ossai, N. I., Nweke, E. O., Ogbuokiri, B. O., Eneh, A. H., Ugwuishiwu, C. H., & Christiana, A. "
        "(2022). An internet of things labelled dataset for aquaponics fish pond water quality monitoring system. "
        "Data in Brief, 43, 108400.",
        "Udanor, C., Ogbuokiri, B., & Nweke, O. (2022). Sensor Based Aquaponics Fish Pond Datasets [Dataset]. Kaggle. "
        "https://www.kaggle.com/datasets/ogbuokiriblessing/sensor-based-aquaponics-fish-pond-datasets",
        "Purbaningtyas, R. Metode Perangkingan Weighted Product dan SAW. Materi Kuliah Sistem Pendukung Keputusan.",
        "Jurnal contoh penerapan Weighted Product pada seleksi penerima beasiswa. Jurnal RESTI. "
        "http://www.jurnal.iaii.or.id/index.php/RESTI/article/view/834/128",
        "Jurnal contoh penerapan SAW pada pemilihan siswa-siswi berprestasi. "
        "http://ejurnal.seminar-id.com/index.php/tin/article/view/655/438",
    ]
    for i, r in enumerate(refs, 1):
        p = rp.item(f"{i}.", r, left=0.9, hang=0.9, after=6)
        p.paragraph_format.alignment = L
    rp.P("Catatan: lengkapi data jurnal (penulis, judul, tahun, volume) dari dua tautan jurnal contoh sesuai "
         "format sitasi program studi Anda.", hl=True, first=False)


# ------------------------------------------------------------------ pasca-proses
def set_page_numbers(doc):
    secs = doc.sections
    # sec0 = cover (tanpa nomor), sec1 = romawi, sec2 = angka
    for idx, fmt in ((1, "lowerRoman"), (2, "decimal")):
        s = secs[idx]
        s.footer.is_linked_to_previous = False
        sectPr = s._sectPr
        pg = sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            mar = sectPr.find(qn("w:pgMar"))
            mar.addnext(pg)
        pg.set(qn("w:fmt"), fmt)
        pg.set(qn("w:start"), "1")
        p = s.footer.paragraphs[0]
        p.paragraph_format.alignment = C
        for kind, txt in (("begin", None), ("instr", "PAGE"), ("end", None)):
            r = p.add_run()
            r.font.name = FONT
            r.font.size = Pt(12)
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
    secs[0].footer.is_linked_to_previous = False


def isi_toc(rp, toc_pages):
    """Sisipkan Daftar Isi sebagai field TOC (dengan hasil awal) di tempat anchor."""
    d = rp.doc
    anchor = rp.toc_anchor
    entries = rp.headings
    paras = []
    for i, (lvl, text) in enumerate(entries):
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(3)
        pf.line_spacing = 1.15
        pf.left_indent = Cm(0 if lvl == 1 else (0.6 if lvl == 2 else 1.2))
        pf.tab_stops.add_tab_stop(Cm(14), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        pg = toc_pages.get(text, "") if toc_pages else ""
        add_runs(p, f"{text}\t{pg}", bold=(lvl == 1))
        paras.append(p)

    def fld(p, kind, instr=None, front=False):
        r = OxmlElement("w:r")
        if instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r.append(el)
        if front:
            p._p.insert(1 if p._p.pPr is not None else 0, r)
        else:
            p._p.append(r)
        return r

    first, last = paras[0], paras[-1]
    # begin, instr, separate harus berurutan di depan paragraf pertama
    pos = 1 if first._p.pPr is not None else 0
    for k, (kind, ins) in enumerate((("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None))):
        r = OxmlElement("w:r")
        if ins:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = ins
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r.append(el)
        first._p.insert(pos + k, r)
    fld(last, "end")
    # pindahkan semua paragraf TOC ke posisi anchor
    for p in paras:
        anchor._p.addprevious(p._p)
    anchor._p.getparent().remove(anchor._p)
    # perbaiki atribut zoom bawaan template agar lolos validasi skema
    z = d.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None:
        z.set(qn("w:percent"), "100")


def halaman_heading(pdf_path, headings):
    """Cari nomor halaman (nomor body, mulai 1 di BAB I) tiap heading dari PDF hasil render."""
    out = subprocess.run(["pdftotext", "-layout", pdf_path, "-"], capture_output=True, text=True).stdout
    pages = out.split("\f")
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    start = None
    for i, pg in enumerate(pages):
        lines = [norm(x) for x in pg.splitlines()]
        if any(l.startswith("BAB I PROFIL DATASET") or l == "BAB I PROFIL DATASET" for l in lines):
            # abaikan halaman daftar isi (memuat titik-titik)
            if not any("....." in l for l in lines):
                start = i
                break
    res = {}
    if start is None:
        return res
    cur = start
    for lvl, text in headings:
        key = norm(text)[:45]
        for i in range(cur, len(pages)):
            lines = [norm(x) for x in pages[i].splitlines()]
            if any(l.startswith(key) for l in lines):
                res[text] = str(i - start + 1)
                cur = i
                break
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", help="folder berisi berkas CSV Kaggle (satu berkas = satu kolam)")
    ap.add_argument("--n", type=int, default=None, help="jumlah kolam yang dipakai")
    ap.add_argument("--out", default="Laporan_SPK_Perangkingan_Tambak.docx")
    a = ap.parse_args()

    if a.data:
        raw, refs = core.load_folder(a.data, a.n)
        if raw.empty:
            raise SystemExit("Tidak ada kolom suhu/kekeruhan/pH/oksigen yang terdeteksi. Periksa nama kolom CSV.")
        demo = False
    else:
        raw, refs, demo = core.data_demo(), None, True
    h = core.hitung_semua(raw, refs=refs)

    with tempfile.TemporaryDirectory() as tmp:
        # pass 1: tanpa nomor halaman
        rp = Rpt()
        bangun(rp, h, demo, tmp)
        set_page_numbers(rp.doc)
        isi_toc(rp, None)
        p1 = os.path.join(tmp, "pass1.docx")
        rp.doc.save(p1)
        pages = {}
        try:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, p1],
                           capture_output=True, timeout=180)
            pages = halaman_heading(os.path.join(tmp, "pass1.pdf"), rp.headings)
        except Exception as e:  # LibreOffice tidak ada -> TOC tanpa nomor, Word akan memperbarui
            print("Peringatan: nomor halaman daftar isi tidak dihitung:", e)
        # pass 2: isi nomor halaman
        rp = Rpt()
        bangun(rp, h, demo, tmp)
        set_page_numbers(rp.doc)
        isi_toc(rp, pages)
        rp.doc.save(a.out)
    print("Laporan tersimpan:", a.out)


if __name__ == "__main__":
    main()
